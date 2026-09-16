#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
雲端版教練總覽生成器（GitHub Actions 用）
從 Dropbox API 讀 xlsx，生成 HTML，再上傳回 Dropbox。
"""
import os
from io import BytesIO
from datetime import datetime, timezone, timedelta
TZ_TPE = timezone(timedelta(hours=8))
import dropbox
from openpyxl import load_workbook

from dropbox_auth import get_dbx
BASE = "/教練業務管理"
OUTPUT = f"{BASE}/教練文件/教練總覽.html"

STUDENTS = [
    {"code": "M26001", "name": "王奕翔", "project": "馬拉松",
     "health": f"{BASE}/學員專用/馬拉松學員/M26001-王奕翔/M26001-王奕翔_健康紀錄.xlsx"},
    {"code": "M26002", "name": "林洋樂", "project": "馬拉松",
     "health": f"{BASE}/學員專用/馬拉松學員/M26002-林洋樂/M26002-林洋樂_健康紀錄.xlsx"},
    {"code": "M26003", "name": "張維倫", "project": "馬拉松",
     "health": f"{BASE}/學員專用/馬拉松學員/M26003-張維倫/M26003-張維倫_健康紀錄.xlsx"},
    {"code": "T26001", "name": "阮筱軒", "project": "國考",
     "health": f"{BASE}/學員專用/國考體測學員/T26001-阮筱軒/T26001-阮筱軒_健康紀錄.xlsx"},
    {"code": "T26002", "name": "盧冠婷", "project": "國考",
     "health": f"{BASE}/學員專用/國考體測學員/T26002-盧冠婷/T26002-盧冠婷_健康紀錄.xlsx"},
]

_db = None
def dbx():
    global _db
    if _db is None:
        _db = get_dbx()
    return _db

def read_xlsx(path):
    try:
        _, res = dbx().files_download(path)
        return load_workbook(BytesIO(res.content), data_only=True)
    except Exception as e:
        print(f"⚠ 讀取失敗 {path}: {e}")
        return None

def read_dashboard(path):
    wb = read_xlsx(path)
    if not wb or "邊界與工作紀錄" not in wb.sheetnames:
        if wb: wb.close()
        return {"boundary": {}, "work": {}}
    ws = wb["邊界與工作紀錄"]
    boundary, work = {}, {}
    section = None
    for row in ws.iter_rows(values_only=True):
        if not any(row): continue
        first = str(row[0]) if row[0] else ""
        if "邊界與約束" in first: section = "boundary"; continue
        if "工作狀態" in first: section = "work"; continue
        if first in ("項目",): continue
        if first.startswith("讀取規定") or "資料讀取" in first: continue
        key = first.strip()
        val = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if not key or key.startswith("王奕翔") or key.startswith("M260") or key.startswith("T260"): continue
        if section == "boundary": boundary[key] = val
        elif section == "work": work[key] = val
    wb.close()
    return {"boundary": boundary, "work": work}

def read_active_injuries(path):
    """讀健康紀錄傷病史，回傳目前恢復中的傷"""
    wb = read_xlsx(path)
    if not wb or "傷病史" not in wb.sheetnames:
        if wb: wb.close()
        return []
    ws = wb["傷病史"]
    injuries = []
    for r in range(2, ws.max_row+1):
        date_v = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        loc = ws.cell(r, 3).value
        status = ws.cell(r, 7).value
        if not name: continue
        status_str = str(status) if status else ""
        # 只顯示未痊癒的
        if "已痊癒" in status_str: continue
        text = f"{name}"
        if loc: text += f"（{loc}）"
        if status_str: text += f"，{status_str}"
        if date_v: text = f"[{date_v}] " + text
        injuries.append(text)
    wb.close()
    return injuries

def read_today_schedule(student):
    code, project = student["code"], student["project"]
    if project == "馬拉松":
        plan_path = f"{BASE}/學員專用/馬拉松學員/{code}-{student['name']}/{code}-{student['name']}_馬拉松課表.xlsx"
        class_path = f"{BASE}/教練專用學員資料/馬拉松/{code}-{student['name']}/{code}-{student['name']}_實體課_教練專用.xlsx"
    else:
        plan_path = f"{BASE}/學員專用/國考體測學員/{code}-{student['name']}/{code}-{student['name']}_國考課表.xlsx"
        class_path = f"{BASE}/教練專用學員資料/國考體測/{code}-{student['name']}/{code}-{student['name']}_實體課_教練專用.xlsx"

    today = datetime.now(TZ_TPE)
    patterns = [today.strftime("%m/%d"), f"{today.month}/{today.day}",
                today.strftime("%Y-%m-%d"), f"{today.year}/{today.month}/{today.day}"]
    def match_date(s):
        s = str(s)
        return any(p in s for p in patterns)

    wb = read_xlsx(plan_path)
    if not wb: return None
    content = note = None
    content_col = note_col = None
    for sn in wb.sheetnames:
        if "每日" in sn:
            ws = wb[sn]
            for header_row in ws.iter_rows(values_only=True):
                if header_row and any("日期" in str(c) for c in header_row if c):
                    for idx, cell in enumerate(header_row):
                        cs = str(cell) if cell else ""
                        if cs in ("訓練內容", "內容"): content_col = idx
                        if "備註" in cs: note_col = idx
                    break
            if content_col is None: content_col = 2
            if note_col is None: note_col = 3
            for row in ws.iter_rows(values_only=True):
                if row and row[0] and match_date(row[0]):
                    content = str(row[content_col]) if len(row) > content_col and row[content_col] else ""
                    note = str(row[note_col]) if len(row) > note_col and row[note_col] else ""
                    break
            break
    wb.close()
    if not content: return None

    result = {"content": content, "note": note, "is_class": "實體課" in content, "steps": []}
    result["cancelled"] = ("❌" in note or "取消" in note or "請假" in note)

    if result["is_class"]:
        cwb = read_xlsx(class_path)
        if cwb:
            for csn in cwb.sheetnames:
                cws = cwb[csn]
                rows = list(cws.iter_rows(values_only=True))
                for i, row in enumerate(rows):
                    if row[1] and match_date(row[1]):
                        steps = []
                        j = i
                        while j < len(rows):
                            r = rows[j]
                            if j > i and r[0] and str(r[0]).strip() and str(r[0]).strip() not in ("堂次",):
                                break
                            if r[3] and str(r[3]).strip():
                                steps.append((str(r[3]).strip(), str(r[4]).strip() if r[4] else "",
                                              str(r[5]).strip() if len(r) > 5 and r[5] else ""))
                            j += 1
                        result["steps"] = steps
                        break
                break
            cwb.close()
    return result

def esc(s):
    if not s: return ""
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def render_student(s, data, injuries):
    b, w = data["boundary"], data["work"]
    tag_class = "tag-green" if s["project"] == "國考" else "tag"
    rows = []
    for k in ["目標", "B 級賽事", "教練出國", "固定不能訓練", "實體課時間", "手錶"]:
        if k in b and b[k]:
            rows.append(f"<tr><td>{k}</td><td>{esc(b[k])}</td></tr>")
    if injuries:
        inj_html = "<br>".join(f'<span style="color:#e74c3c">🩹 {esc(i)}</span>' for i in injuries)
        rows.append(f"<tr><td>目前傷病</td><td>{inj_html}</td></tr>")
    for k in ["目前課表版本", "最近調整", "下次待辦", "待決定"]:
        if k in w and w[k]:
            rows.append(f"<tr><td>{k}</td><td>{esc(w[k])}</td></tr>")
    if not rows:
        rows.append('<tr><td colspan="2" style="color:#999">待補資料</td></tr>')
    return f"""
<div class="student">
  <h3>👤 {s['code']} {s['name']} <span class="{tag_class}">{s['project']}</span></h3>
  <table>{''.join(rows)}</table>
</div>"""

def main():
    now = datetime.now(TZ_TPE).strftime("%Y-%m-%d %H:%M")
    weekday_cn = "一二三四五六日"[datetime.now(TZ_TPE).weekday()]
    students_html, todos, today_sessions = [], [], []

    for s in STUDENTS:
        data = read_dashboard(s["health"])
        injuries = read_active_injuries(s["health"])
        students_html.append(render_student(s, data, injuries))
        todo = data["work"].get("下次待辦", "")
        if todo:
            todos.append(f"<li><strong>{s['code']} {s['name']}</strong>：{esc(todo)}</li>")
        # 傷病也進待辦提醒
        for inj in injuries:
            todos.append(f'<li style="color:#e74c3c"><strong>{s["code"]} {s["name"]}</strong>：傷病追蹤 - {esc(inj)}</li>')

        t = read_today_schedule(s)
        recent = data["work"].get("最近調整", "")
        nowd = datetime.now(TZ_TPE)
        today_keys = {f"{nowd.month}/{nowd.day}", nowd.strftime("%m/%d"),
                      nowd.strftime("%Y-%m-%d"), nowd.strftime("%m-%d")}
        neg_kw = ["取消", "請假", "順延", "改期", "暫停", "受傷", "因公", "未執行"]
        pos_kw = ["正常進行", "正常上課", "正常上", "已痊癒", "恢復正常", "正常參與"]
        adjustment_today = ""
        if recent:
            items = [ln.strip() for ln in recent.split("\n") if ln.strip()]
            today_items = [ln for ln in items if any(k in ln for k in today_keys)]
            neg_items = [ln for ln in today_items
                         if any(k in ln for k in neg_kw) and not any(p in ln for p in pos_kw)]
            if neg_items:
                adjustment_today = "；".join(neg_items)[:200]

        if t and t["content"]:
            label = "已取消" if (t.get("cancelled") or adjustment_today) else ("實體課" if t["is_class"] else "自主訓練")
            item = f"<li><strong>{s['code']} {s['name']}</strong>（{label}）：{esc(t['content'])}"
            if t["note"]:
                item += f' <span style="color:#888">｜{esc(t["note"])}</span>'
            item += "</li>"
            if adjustment_today:
                item += f'<div style="margin-top:3px;padding-left:15px;color:#e74c3c;font-size:13px">📝 今日異動：{esc(adjustment_today)[:150]}</div>'
            # 傷病提醒
            if injuries:
                inj_text = "；".join(injuries)
                item += f'<div style="margin-top:3px;padding-left:15px;color:#e74c3c;font-size:13px">🩹 傷病：{esc(inj_text)[:200]}</div>'
            if t["steps"] and not t.get("cancelled") and not adjustment_today:
                sh = "<ul style='margin-top:5px'>"
                for sn, tm, sc in t["steps"]:
                    sh += f"<li><strong>{esc(sn)}</strong>（{esc(tm)}）：{esc(sc)}</li>"
                sh += "</ul>"
                item = item.replace("</li>", sh + "</li>")
            today_sessions.append(item)
        elif adjustment_today:
            today_sessions.append(f'<li><strong>{s["code"]} {s["name"]}</strong>：<div style="color:#e74c3c">📝 {esc(adjustment_today)[:200]}</div></li>')

    today_block = ""
    if today_sessions:
        today_block = f"""
<div class="today-block">
  <strong>📅 今日焦點（{datetime.now(TZ_TPE).strftime('%m/%d')} 週{weekday_cn}）</strong>
  <ul>{''.join(today_sessions)}</ul>
</div>"""

    html = f"""<!DOCTYPE html>
<html lang="zh-Hant"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>教練總覽</title>
<style>
  body {{ font-family: -apple-system, "PingFang TC", "Microsoft JhengHei", sans-serif; max-width: 900px; margin: 20px auto; padding: 0 15px; color: #333; line-height: 1.6; }}
  h1 {{ color: #2c3e50; border-bottom: 3px solid #4472C4; padding-bottom: 10px; font-size: 24px; }}
  .updated {{ color: #888; font-size: 13px; margin-bottom: 15px; }}
  .todo {{ background: #fff3cd; padding: 12px 18px; border-radius: 8px; margin: 12px 0; }}
  .todo ol, .today-block ul {{ margin: 5px 0; padding-left: 22px; }}
  .today-block {{ background: #d4edda; padding: 12px 18px; border-radius: 8px; margin: 12px 0; border-left: 5px solid #27ae60; }}
  table {{ border-collapse: collapse; width: 100%; margin: 8px 0; font-size: 14px; }}
  th, td {{ border: 1px solid #ddd; padding: 6px 10px; text-align: left; }}
  th {{ background: #4472C4; color: white; }}
  tr:nth-child(even) {{ background: #f9f9f9; }}
  .student {{ background: #f0f4fa; padding: 12px 18px; border-radius: 8px; margin: 12px 0; }}
  .student h3 {{ margin-top: 0; color: #2c3e50; font-size: 17px; }}
  .tag {{ display: inline-block; background: #4472C4; color: white; padding: 2px 8px; border-radius: 4px; font-size: 12px; }}
  .tag-green {{ background: #27ae60; }}
</style></head><body>
<h1>🏃 教練總覽</h1>
<div class="updated">自動更新時間：{now}</div>
{today_block}
<div class="todo"><strong>📌 目前待辦</strong><ol>{''.join(todos) if todos else '<li>無</li>'}</ol></div>
{''.join(students_html)}
</body></html>"""

    # 上傳到 Dropbox
    dbx().files_upload(
        html.encode("utf-8"),
        OUTPUT,
        mode=dropbox.files.WriteMode("overwrite")
    )
    print(f"✓ 已上傳到 Dropbox：{OUTPUT}")

if __name__ == "__main__":
    main()
