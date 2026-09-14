#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
雲端版晨間簡報（GitHub Actions 用）
從 Dropbox API 讀 xlsx，用 SMTP 發 email。
"""
import os
import re
import smtplib
from io import BytesIO
from datetime import datetime, timedelta
from email.mime.text import MIMEText
from email.header import Header

import dropbox
from openpyxl import load_workbook

# === 設定 ===
DROPBOX_TOKEN = os.environ["DROPBOX_TOKEN"]
SMTP_USER = os.environ["SMTP_USER"]
SMTP_PASS = os.environ["SMTP_PASS"]
COACH_EMAIL = "runningcoach.kevin@gmail.com"
BASE = "/教練業務管理"

CIVIL_DIR = f"{BASE}/學員專用/國考體測學員"
MARATHON_DIR = f"{BASE}/學員專用/馬拉松學員"
INPERSON_CIVIL = f"{BASE}/教練專用學員資料/國考體測"
INPERSON_MARATHON = f"{BASE}/教練專用學員資料/馬拉松"

STUDENTS = [
    {"編號": "T26001", "姓名": "阮筱軒", "專案": "國考",
     "日誌": f"{CIVIL_DIR}/T26001-阮筱軒/T26001-阮筱軒_國考訓練日誌.xlsx",
     "課表": f"{CIVIL_DIR}/T26001-阮筱軒/T26001-阮筱軒_國考課表.xlsx"},
    {"編號": "T26002", "姓名": "盧冠婷", "專案": "國考",
     "日誌": f"{CIVIL_DIR}/T26002-盧冠婷/T26002-盧冠婷_國考訓練日誌.xlsx",
     "課表": f"{CIVIL_DIR}/T26002-盧冠婷/T26002-盧冠婷_國考課表.xlsx"},
    {"編號": "M26001", "姓名": "王奕翔", "專案": "馬拉松",
     "日誌": f"{MARATHON_DIR}/M26001-王奕翔/M26001-王奕翔_馬拉松訓練日誌.xlsx",
     "課表": f"{MARATHON_DIR}/M26001-王奕翔/M26001-王奕翔_馬拉松課表.xlsx"},
    {"編號": "M26002", "姓名": "林洋樂", "專案": "馬拉松",
     "日誌": f"{MARATHON_DIR}/M26002-林洋樂/M26002-林洋樂_馬拉松訓練日誌.xlsx",
     "課表": f"{MARATHON_DIR}/M26002-林洋樂/M26002-林洋樂_馬拉松課表.xlsx"},
]

WEEKDAY_MAP = {0: "一", 1: "二", 2: "三", 3: "四", 4: "五", 5: "六", 6: "日"}

# === Dropbox 讀檔 ===
_db = None
def dbx():
    global _db
    if _db is None:
        _db = dropbox.Dropbox(DROPBOX_TOKEN)
    return _db

def read_xlsx_from_dropbox(path):
    """從 Dropbox 讀 xlsx，回傳 workbook 物件。"""
    try:
        metadata, res = dbx().files_download(path)
        return load_workbook(BytesIO(res.content), data_only=True)
    except Exception as e:
        print(f"⚠ 讀取失敗 {path}: {e}")
        return None

# === 讀課表 ===
def read_schedule(schedule_path, target_date):
    wb = read_xlsx_from_dropbox(schedule_path)
    if not wb or "每日課表" not in wb.sheetnames:
        if wb: wb.close()
        return None
    ws = wb["每日課表"]
    patterns = [
        target_date.strftime("%-m/%-d"),
        target_date.strftime("%m/%d"),
        target_date.strftime("%Y-%m-%d"),
    ]
    def match_date(v):
        if not v: return False
        v = str(v)
        return any(p in v for p in patterns)
    content_col, note_col = 4, None
    for r in range(1, min(ws.max_row+1, 10)):
        vals = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        if any("日期" in str(v) for v in vals if v):
            for c, v in enumerate(vals, 1):
                vs = str(v) if v else ""
                if vs in ("訓練內容", "內容"): content_col = c
                if "備註" in vs: note_col = c
            break
    for r in range(1, ws.max_row+1):
        for c in range(1, min(ws.max_column+1, 5)):
            if match_date(ws.cell(r,c).value):
                content = ws.cell(r, content_col).value
                note = ""
                if note_col:
                    n = ws.cell(r, note_col).value
                    note = str(n).strip() if n else ""
                result = str(content).strip() if content else "無紀錄"
                if note: result += f"（{note}）"
                wb.close()
                return result
    wb.close()
    return None

# === 讀日誌 ===
def read_log_summary(log_path, target_date):
    wb = read_xlsx_from_dropbox(log_path)
    if not wb:
        return None
    # 找訓練記錄分頁
    sheet_name = None
    for sn in wb.sheetnames:
        if "訓練記錄" in sn or "日誌" in sn or "每日" in sn:
            sheet_name = sn; break
    if not sheet_name:
        wb.close(); return None
    ws = wb[sheet_name]
    patterns = [target_date.strftime("%-m/%-d"), target_date.strftime("%Y-%m-%d")]
    for r in range(1, ws.max_row+1):
        for c in range(1, min(ws.max_column+1, 5)):
            v = ws.cell(r,c).value
            if v and any(p in str(v) for p in patterns):
                # 讀內容欄
                for cc in range(c+1, min(ws.max_column+1, 15)):
                    val = ws.cell(r, cc).value
                    if val and len(str(val)) > 3:
                        wb.close()
                        return str(val).strip()[:80]
    wb.close()
    return None

# === 讀下一次實體課 ===
def read_next_inperson(student_id, student_name, today):
    project_dir = INPERSON_CIVIL if student_id.startswith("T") else INPERSON_MARATHON
    path = f"{project_dir}/{student_id}-{student_name}/{student_id}-{student_name}_實體課_教練專用.xlsx"
    wb = read_xlsx_from_dropbox(path)
    if not wb: return None
    ws = wb["實體課"] if "實體課" in wb.sheetnames else wb.active
    for r in range(3, ws.max_row+1):
        date_raw = ws.cell(r, 2).value
        topic = ws.cell(r, 3).value
        tang = ws.cell(r, 1).value
        if date_raw and topic and tang and tang != "—":
            m = re.match(r'(\d+)/(\d+)', str(date_raw))
            if m:
                month, day = int(m.group(1)), int(m.group(2))
                session_date = datetime(today.year, month, day).date()
                if session_date >= today:
                    wb.close()
                    return f"{tang}｜{month}/{day}｜{str(topic).strip()}"
    wb.close()
    return None

# === 組裝 ===
def generate_briefing(target_date):
    yesterday = target_date - timedelta(days=1)
    tomorrow = target_date + timedelta(days=1)
    weekday = WEEKDAY_MAP[target_date.weekday()]
    lines = [f"🌅 晨間訓練簡報｜{target_date.strftime('%-m/%-d')} {weekday}", "="*40]

    for proj_name, proj_filter in [("國考體測學員", "國考"), ("馬拉松學員", "馬拉松")]:
        lines.append(f"\n━━━ {proj_name} ━━━")
        lines.append("\n📋 今日課表總覽")
        for s in STUDENTS:
            if s["專案"] != proj_filter: continue
            t = read_schedule(s["課表"], target_date)
            lines.append(f"  {s['編號']}-{s['姓名']}：{t or '無紀錄'}")
        lines.append("\n📊 昨日訓練摘要")
        for s in STUDENTS:
            if s["專案"] != proj_filter: continue
            l = read_log_summary(s["日誌"], yesterday)
            lines.append(f"  {s['編號']}-{s['姓名']}：{l or '休息/無紀錄'}")
        lines.append("\n📅 明日課表預告")
        for s in STUDENTS:
            if s["專案"] != proj_filter: continue
            t = read_schedule(s["課表"], tomorrow)
            lines.append(f"  {s['編號']}-{s['姓名']}：{t or '無紀錄'}")

    lines.append("\n━━━ 🏋️ 下一次實體課預告 ━━━")
    for s in STUDENTS:
        n = read_next_inperson(s["編號"], s["姓名"], target_date)
        lines.append(f"  {s['編號']}-{s['姓名']}：{n or '無待排實體課'}")

    lines.append("\n" + "="*40)
    lines.append("本簡報由訓練通知系統自動產生")
    return "\n".join(lines)

# === 寄信 ===
def send_email(body):
    msg = MIMEText(body, "plain", "utf-8")
    msg["From"] = SMTP_USER
    msg["To"] = COACH_EMAIL
    msg["Subject"] = Header(f"🌅 晨間訓練簡報｜{datetime.now().strftime('%-m/%-m')}", "utf-8")
    with smtplib.SMTP("smtp.gmail.com", 587) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASS)
        server.sendmail(SMTP_USER, [COACH_EMAIL], msg.as_string())
    print(f"✓ 已寄到 {COACH_EMAIL}")

if __name__ == "__main__":
    today = datetime.now().date()
    body = generate_briefing(today)
    print(body)
    send_email(body)
