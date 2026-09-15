#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
雲端版學員每日通知（GitHub Actions 用）
從 Dropbox API 讀 xlsx，寄 email 給每位學員，CC 教練。
"""
import os
import re
import smtplib
from io import BytesIO
from datetime import datetime, timedelta
from email.mime.text import MIMEText
from email.header import Header

import dropbox
from openpyxl import load_workbook

from dropbox_auth import get_dbx
SMTP_USER = os.environ["SMTP_USER"]
SMTP_PASS = os.environ["SMTP_PASS"]
COACH_EMAIL = "runningcoach.kevin@gmail.com"
BASE = "/教練業務管理"

CIVIL_DIR = f"{BASE}/學員專用/國考體測學員"
MARATHON_DIR = f"{BASE}/學員專用/馬拉松學員"

STUDENTS = [
    {"編號": "M26001", "姓名": "王奕翔", "email": "mike141431@gmail.com",
     "日誌": f"{MARATHON_DIR}/M26001-王奕翔/M26001-王奕翔_馬拉松訓練日誌.xlsx",
     "課表": f"{MARATHON_DIR}/M26001-王奕翔/M26001-王奕翔_馬拉松課表.xlsx",
     "健康": f"{MARATHON_DIR}/M26001-王奕翔/M26001-王奕翔_健康紀錄.xlsx"},
    {"編號": "M26002", "姓名": "林洋樂", "email": "alerler0817@gmail.com",
     "日誌": f"{MARATHON_DIR}/M26002-林洋樂/M26002-林洋樂_馬拉松訓練日誌.xlsx",
     "課表": f"{MARATHON_DIR}/M26002-林洋樂/M26002-林洋樂_馬拉松課表.xlsx",
     "健康": f"{MARATHON_DIR}/M26002-林洋樂/M26002-林洋樂_健康紀錄.xlsx"},
    {"編號": "T26001", "姓名": "阮筱軒", "email": "syuan000906@gmail.com",
     "日誌": f"{CIVIL_DIR}/T26001-阮筱軒/T26001-阮筱軒_國考訓練日誌.xlsx",
     "課表": f"{CIVIL_DIR}/T26001-阮筱軒/T26001-阮筱軒_國考課表.xlsx",
     "健康": f"{CIVIL_DIR}/T26001-阮筱軒/T26001-阮筱軒_健康紀錄.xlsx"},
    {"編號": "T26002", "姓名": "盧冠婷", "email": "tina19981217@gmail.com",
     "日誌": f"{CIVIL_DIR}/T26002-盧冠婷/T26002-盧冠婷_國考訓練日誌.xlsx",
     "課表": f"{CIVIL_DIR}/T26002-盧冠婷/T26002-盧冠婷_國考課表.xlsx",
     "健康": f"{CIVIL_DIR}/T26002-盧冠婷/T26002-盧冠婷_健康紀錄.xlsx"},
]

WEEKDAY_MAP = {0: "一", 1: "二", 2: "三", 3: "四", 4: "五", 5: "六", 6: "日"}

_db = None
def dbx():
    global _db
    if _db is None:
        _db = get_dbx()
    return _db

def read_xlsx(path):
    try:
        _, res = dbx().files_download(path)
        return load_workbook(BytesIO(res.content), data_only=True)
    except Exception as e:
        print(f"⚠ 讀取失敗 {path}: {e}")
        return None

def read_schedule(path, target_date):
    wb = read_xlsx(path)
    if not wb or "每日課表" not in wb.sheetnames:
        if wb: wb.close()
        return None
    ws = wb["每日課表"]
    patterns = [
        target_date.strftime("%-m/%-d"),
        target_date.strftime("%m/%d"),
        target_date.strftime("%Y-%m-%d"),
    ]
    content_col = 4
    for r in range(1, min(ws.max_row+1, 10)):
        vals = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        if any("日期" in str(v) for v in vals if v):
            for c, v in enumerate(vals, 1):
                if str(v) in ("訓練內容", "內容"):
                    content_col = c
            break
    for r in range(1, ws.max_row+1):
        for c in range(1, min(ws.max_column+1, 5)):
            v = ws.cell(r,c).value
            if v and any(p in str(v) for p in patterns):
                content = ws.cell(r, content_col).value
                wb.close()
                return str(content).strip() if content else None
    wb.close()
    return None

def read_log(path, target_date):
    wb = read_xlsx(path)
    if not wb: return None
    ws = wb.active
    patterns = [target_date.strftime("%-m/%-d"), target_date.strftime("%Y-%m-%d")]
    for r in range(1, ws.max_row+1):
        v = ws.cell(r, 1).value
        if v and any(p in str(v) for p in patterns):
            # C=課表內容 D=距離 E=時間 F=平均心率 G=最高心率 H=教練評註
            record = {
                "距離": ws.cell(r, 4).value,
                "時間": ws.cell(r, 5).value,
                "平均心率": ws.cell(r, 6).value,
                "最高心率": ws.cell(r, 7).value,
                "教練評註": ws.cell(r, 8).value,
            }
            wb.close()
            return record
    wb.close()
    return None

def read_health(path, target_date):
    """讀健康紀錄，回傳昨天的受傷/異常狀態"""
    wb = read_xlsx(path)
    if not wb: return None
    result = {"injury": None, "notes": []}

    patterns = [target_date.strftime("%-m/%-d"), target_date.strftime("%Y-%m-%d"),
                target_date.strftime("%Y/%-m/%-d")]

    # 傷病史
    if "傷病史" in wb.sheetnames:
        ws = wb["傷病史"]
        for r in range(2, ws.max_row+1):
            date_v = ws.cell(r, 1).value
            if date_v and any(p in str(date_v) for p in patterns):
                injury_name = ws.cell(r, 2).value
                location = ws.cell(r, 3).value
                status = ws.cell(r, 7).value
                if injury_name:
                    result["injury"] = f"{injury_name}"
                    if location: result["injury"] += f"（{location}）"
                    if status: result["injury"] += f"，{status}"

    # 每日監測：主觀感覺差
    if "每日監測" in wb.sheetnames:
        ws = wb["每日監測"]
        for r in range(2, ws.max_row+1):
            date_v = ws.cell(r, 1).value
            if date_v and any(p in str(date_v) for p in patterns):
                # 找主觀感覺欄
                feel = None
                for c in range(1, ws.max_column+1):
                    h = ws.cell(1, c).value
                    if h and "主觀" in str(h):
                        feel = ws.cell(r, c).value
                        break
                if feel and isinstance(feel, (int, float)) and feel <= 2:
                    result["notes"].append(f"昨日主觀感覺 {feel}/5，較差")
                break

    wb.close()
    return result if (result["injury"] or result["notes"]) else None

def get_first_name(full_name):
    if not full_name or len(full_name) <= 1:
        return full_name
    return full_name[1:]

def format_content(content):
    if not content: return content
    content = str(content).strip()
    # 實體課：只保留「實體課」+ 時間地點，不揭露操課細節
    if "實體課" in content:
        # 抓括號裡的時間地點
        loc_match = re.search(r'[（(]([^）)]*)[）)]', content)
        loc = loc_match.group(1) if loc_match else ""
        # 抓堂次
        session_match = re.search(r'實體課[①②③④⑤⑥⑦⑧⑨⑩\d]*', content)
        session = session_match.group(0) if session_match else "實體課"
        if loc:
            return f"{session}（{loc}）"
        return session
    # 一般訓練：保留配速，不刪
    return content

def validate_email(today_s, tomorrow_s, yest, yest_plan_str, today):
    """AI 自動檢查，回傳 warning list"""
    warnings = []

    # 1. 實體課細節不該露出（操課術語）
    drill_keywords = ["×", "組間", "recovery", "Recovery", "m@", "m @", "/圈"]
    for label, content in [("今天", today_s), ("明天", tomorrow_s)]:
        if "實體課" in content:
            for kw in drill_keywords:
                if kw in content:
                    warnings.append(f"⚠️ {label}課表疑似露出實體課操課細節（含「{kw}」）")
                    break

    # 2. 受傷後隔天還排高強度
    if "受傷" in yest_plan_str:
        high_intensity = any(kw in today_s + tomorrow_s for kw in ["間歇", "T跑", "I跑", "5K", "測驗", "TT", "模擬"])
        if high_intensity:
            warnings.append("⚠️ 昨天受傷，今天/明天排了高強度課，確認是否該改休息")

    # 3. 空值 / None / NA
    for label, content in [("今天", today_s), ("明天", tomorrow_s), ("昨天", yest)]:
        if content is None or str(content).strip() in ("", "None", "NA", "nan"):
            warnings.append(f"⚠️ {label}課表內容是空的")

    # 4. 日期星期檢查
    weekday = today.weekday()
    weekday_cn = WEEKDAY_MAP[weekday]
    # 檢查今天內容裡的星期標籤是否正確
    expected = f"（{weekday_cn}）"
    # 這個比較難直接檢，跳過

    return warnings


def build_email(student, today):
    name = student["姓名"]
    first = get_first_name(name)
    yesterday = today - timedelta(days=1)
    tomorrow = today + timedelta(days=1)

    today_s = read_schedule(student["課表"], today) or "休息日或輕鬆恢復"
    today_s = format_content(today_s)

    tomorrow_s = read_schedule(student["課表"], tomorrow) or "休息日或輕鬆恢復"
    tomorrow_s = format_content(tomorrow_s)

    log = read_log(student["日誌"], yesterday)
    health = read_health(student["健康"], yesterday)
    day_before_yesterday = today - timedelta(days=2)
    log_prev = read_log(student["日誌"], day_before_yesterday)
    # 先看昨天課表是什麼
    yest_plan = read_schedule(student["課表"], yesterday) or ""
    yest_plan_str = str(yest_plan).strip()
    is_rest_day = ("休息" in yest_plan_str) or (yest_plan_str == "")
    is_leave_day = any(kw in yest_plan_str for kw in ["請假", "取消", "暫停", "因公", "受傷", "順延"])

    # 昨天有受傷記錄（從健康紀錄）
    yesterday_injury = health["injury"] if health else None
    yesterday_health_notes = health["notes"] if health else []

    # 如果昨天受傷，明天課表加備註
    tomorrow_s_display = tomorrow_s
    if yesterday_injury:
        tomorrow_s_display = tomorrow_s + "（視傷勢復原狀況決定是否執行）"

    if log and log.get("距離"):
        yest_plan_fmt = format_content(yest_plan_str) if yest_plan_str else ""
        parts = [f"距離 {log['距離']} km"]
        if log.get("時間"): parts.append(f"時間 {log['時間']}")
        if log.get("平均心率"): parts.append(f"心率 {log['平均心率']}")
        yest = "｜".join(parts)
        if yest_plan_fmt:
            yest = f"昨日課表：{yest_plan_fmt}\n" + yest
        if log.get("教練評註"): yest += f"\n教練評註：{log['教練評註']}"
    elif yesterday_injury:
        yest = f"昨日受傷：{yesterday_injury}"
        yest += "\n好好休息恢復，不要勉強上場；恢復狀況隨時回報給我。"
    elif is_leave_day:
        yest = "昨日請假未訓練"
        if "受傷" in yest_plan_str:
            yest += "\n好好休息恢復，不要勉強上場；恢復狀況隨時回報給我。"
    elif is_rest_day:
        yest = "昨日休息日"
    else:
        yest = "缺昨日訓練數據\n如尚未回傳，請盡速補上"

    # 加上健康備註
    for note in yesterday_health_notes:
        yest += f"\n（{note}）"

    # 前日訓練（前天）
    prev_section = ""
    if log_prev and log_prev.get("教練評註"):
        prev_plan = read_schedule(student["課表"], day_before_yesterday) or ""
        prev_plan_fmt = format_content(prev_plan) if prev_plan else ""
        prev_parts = []
        if log_prev.get("距離"): prev_parts.append(f"距離 {log_prev['距離']} km")
        if log_prev.get("時間"): prev_parts.append(f"時間 {log_prev['時間']}")
        if log_prev.get("平均心率"): prev_parts.append(f"心率 {log_prev['平均心率']}")
        prev_line = "｜".join(prev_parts) if prev_parts else ""
        prev_lines = []
        if prev_plan_fmt: prev_lines.append(f"前日課表：{prev_plan_fmt}")
        if prev_line: prev_lines.append(prev_line)
        if log_prev.get("教練評註"): prev_lines.append(f"教練評註：{log_prev['教練評註']}")
        prev_section = "\n".join(prev_lines)

    weekday = WEEKDAY_MAP[today.weekday()]
    subject = f"早安{first}，這是你的今日學員晨報"

    body = f"""早安，{first}：

很抱歉，近日因系統重整，如有收到錯誤內容的郵件，敬請見諒。

【今天 {today.strftime('%-m/%-d')}（{WEEKDAY_MAP[today.weekday()]}）】
{today_s}

【昨日 {yesterday.strftime('%-m/%-d')}】
{yest}
{f'''
【前日 {day_before_yesterday.strftime('%-m/%-d')}】
{prev_section}''' if prev_section else ''}

【明天 {tomorrow.strftime('%-m/%-d')}（{WEEKDAY_MAP[tomorrow.weekday()]}）】
{tomorrow_s_display}

提醒：
1. 課表內容請照表執行，有問題隨時跟我說
2. 訓練後請用手機截圖資料畫面回傳
3. 身體不適或有任何狀況，請主動回報
4. 課表內容請勿隨意更換或改課，如有需要改動，請回報教練調整

加油！

執行教練｜Kevin Chang
0917060888｜Line: kc1225888
"""
    warnings = validate_email(today_s, tomorrow_s, yest, yest_plan_str, today)

    return subject, body, warnings

DRY_RUN = "--dry-run" in os.environ.get("ARGS", "")

def send(to, subject, body):
    msg = MIMEText(body, "plain", "utf-8")
    msg["From"] = SMTP_USER
    msg["To"] = to
    msg["Cc"] = COACH_EMAIL
    msg["Subject"] = Header(subject, "utf-8")
    if DRY_RUN:
        # 測試模式：寄給教練自己
        print(f"[DRY-RUN] 學員：{to} → 寄給教練 {COACH_EMAIL}")
        test_msg = MIMEText(body, "plain", "utf-8")
        test_msg["From"] = SMTP_USER
        test_msg["To"] = COACH_EMAIL
        test_msg["Subject"] = Header(f"[測試-{to}] {subject}", "utf-8")
        with smtplib.SMTP("smtp.gmail.com", 587) as s:
            s.starttls()
            s.login(SMTP_USER, SMTP_PASS)
            s.sendmail(SMTP_USER, [COACH_EMAIL], test_msg.as_string())
        print(f"  ✓ 已寄到教練信箱")
        return
    with smtplib.SMTP("smtp.gmail.com", 587) as s:
        s.starttls()
        s.login(SMTP_USER, SMTP_PASS)
        s.sendmail(SMTP_USER, [to, COACH_EMAIL], msg.as_string())
    print(f"✓ → {to} (CC {COACH_EMAIL})")

if __name__ == "__main__":
    args = os.environ.get("ARGS", "")
    is_preview = "--preview" in args
    today = datetime.now().date()

    if is_preview:
        # 預審模式：生成明天的信，寄給教練
        tomorrow = today + timedelta(days=1)
        print(f"=== 預審：明天 {tomorrow} 的信件 ===")
        for st in STUDENTS:
            try:
                subject, body, warnings = build_email(st, tomorrow)
                # 如果有 warning，加在信件最前面
                if warnings:
                    warn_block = "【⚠️ AI 自動檢查提醒】\n" + "\n".join(warnings) + "\n\n"
                    body = warn_block + body
                # 寄給教練，主旨加前綴
                test_msg = MIMEText(body, "plain", "utf-8")
                test_msg["From"] = SMTP_USER
                test_msg["To"] = COACH_EMAIL
                test_msg["Subject"] = Header(
                    f"【待審批-{tomorrow.strftime('%m/%d')}】{subject}", "utf-8"
                )
                with smtplib.SMTP("smtp.gmail.com", 587) as s:
                    s.starttls()
                    s.login(SMTP_USER, SMTP_PASS)
                    s.sendmail(SMTP_USER, [COACH_EMAIL], test_msg.as_string())
                status = f"⚠️ {len(warnings)} 個提醒" if warnings else "✓ 無提醒"
                print(f"  ✓ 已寄：{subject}（{status}）")
            except Exception as e:
                print(f"  ✗ {st['姓名']}: {e}")
    else:
        # 正常模式：生成今天的信，寄給學員
        for st in STUDENTS:
            try:
                subject, body, warnings = build_email(st, today)
                send(st["email"], subject, body)
            except Exception as e:
                print(f"✗ {st['姓名']}: {e}")
